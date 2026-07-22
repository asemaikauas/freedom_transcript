#!/usr/bin/env python3
"""Convert the Freedom AI Labs ground-truth workbook to benchmark CSV."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path


DEFAULT_OUTPUT = Path("data/freedom_ai_labs_benchmark.csv")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True, help="Source .xlsx workbook.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--sheet", help="Worksheet name (default: active worksheet).")
    parser.add_argument("--reference-column", default="transcript")
    parser.add_argument("--prediction-column", default="transcript_pred")
    parser.add_argument(
        "--language",
        default="mix",
        choices=("en", "kk", "ru", "mix"),
        help="Language label applied to every row (default: mix).",
    )
    parser.add_argument("--asr-model", default="provided_asr_baseline")
    parser.add_argument("--id-prefix", default="freedom_ai_labs")
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


def load_source_rows(
    input_path: Path,
    sheet_name: str | None,
    reference_column: str,
    prediction_column: str,
) -> list[tuple[int, str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit(
            "openpyxl is required to read .xlsx files. Install requirements.txt first."
        ) from exc

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise SystemExit(
                f"Worksheet {sheet_name!r} was not found. Available: {', '.join(workbook.sheetnames)}"
            )
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.active

    iterator = sheet.iter_rows(values_only=True)
    try:
        header_values = next(iterator)
    except StopIteration as exc:
        raise SystemExit("The source worksheet is empty.") from exc

    headers = [str(value).strip() if value is not None else "" for value in header_values]
    if len(headers) != len(set(headers)):
        raise SystemExit("The source worksheet contains duplicate column names.")

    missing = [name for name in (reference_column, prediction_column) if name not in headers]
    if missing:
        raise SystemExit(
            f"Missing required column(s): {', '.join(missing)}. Found: {', '.join(headers)}"
        )

    reference_index = headers.index(reference_column)
    prediction_index = headers.index(prediction_column)
    rows: list[tuple[int, str, str]] = []
    for excel_row, values in enumerate(iterator, start=2):
        reference = values[reference_index] if reference_index < len(values) else None
        prediction = values[prediction_index] if prediction_index < len(values) else None
        if reference is None and prediction is None:
            continue
        if not isinstance(reference, str) or not reference.strip():
            raise SystemExit(f"Row {excel_row}: {reference_column!r} must be non-empty text.")
        if not isinstance(prediction, str) or not prediction.strip():
            raise SystemExit(f"Row {excel_row}: {prediction_column!r} must be non-empty text.")
        rows.append((excel_row, reference.strip(), prediction.strip()))

    if not rows:
        raise SystemExit("No benchmark rows were found in the source worksheet.")
    return rows


def main() -> None:
    args = parse_args()
    if not args.input.exists():
        raise SystemExit(f"Input workbook does not exist: {args.input}")
    if args.output.exists() and not args.overwrite:
        raise SystemExit(
            f"Refusing to overwrite existing benchmark: {args.output}\n"
            "Pass --overwrite to replace it explicitly."
        )

    rows = load_source_rows(
        args.input,
        args.sheet,
        args.reference_column,
        args.prediction_column,
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "id",
        "example_id",
        "language",
        "asr_model",
        "source_file",
        "row_index",
        "raw_transcript",
        "reference_clean",
        "raw_chars",
        "reference_chars",
    ]
    with args.output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        for example_number, (excel_row, reference, prediction) in enumerate(rows, start=1):
            example_id = f"{args.id_prefix}_{example_number:04d}"
            writer.writerow(
                {
                    "id": example_id,
                    "example_id": example_id,
                    "language": args.language,
                    "asr_model": args.asr_model,
                    "source_file": args.input.name,
                    "row_index": excel_row,
                    "raw_transcript": prediction,
                    "reference_clean": reference,
                    "raw_chars": len(prediction),
                    "reference_chars": len(reference),
                }
            )

    print(f"Wrote {len(rows)} rows to {args.output}")
    print(f"Language={args.language}; ASR baseline={args.asr_model}")


if __name__ == "__main__":
    main()
